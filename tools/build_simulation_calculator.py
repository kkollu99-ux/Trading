"""Builds tools/Simulation-Calculator.xlsx - an Excel sheet (live formulas) for planning a simulation:
type an investment and the % hike you want and it gives the target price (as % and as a price for the
admin simulation form), the lots to buy and the profit in currency and %. Uses the platform's own maths:
1 lot = 100 units, margin = lots*100*price/multiplier, fee = 0.1% of margin, P&L = (exit-entry)*lots*100."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

wb = Workbook()
ws = wb.active
ws.title = "Calculator"
NAVY, BLUE_IN, GREEN, GREY = "0B1F3A", "DCE9FF", "E3F6EA", "F3F5F9"
thin = Side(style="thin", color="C9D1E0")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
hdr = Font(bold=True, color="FFFFFF")
inp_font = Font(color="0033CC", bold=True)

def band(row, text, cols="A:D"):
    ws.merge_cells(f"A{row}:D{row}")
    c = ws[f"A{row}"]; c.value = text; c.font = hdr
    c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 20

def line(row, label, value, note="", fmt=None, inp=False, out=False, bold=False):
    ws[f"A{row}"] = label
    ws[f"B{row}"] = value
    ws[f"C{row}"] = note
    for col in "ABC":
        ws[f"{col}{row}"].border = box
    ws[f"A{row}"].font = Font(bold=bold or out)
    b = ws[f"B{row}"]
    b.alignment = Alignment(horizontal="right")
    if fmt: b.number_format = fmt
    if inp:
        b.fill = PatternFill("solid", fgColor=BLUE_IN); b.font = inp_font
    if out:
        b.fill = PatternFill("solid", fgColor=GREEN); b.font = Font(bold=True)
    ws[f"C{row}"].font = Font(italic=True, color="5F6C84")

ws.merge_cells("A1:D1")
ws["A1"] = "FXCM Simulation Calculator"; ws["A1"].font = title
ws["A1"].fill = PatternFill("solid", fgColor=NAVY); ws.row_dimensions[1].height = 30
ws["A2"] = "Blue cells are inputs - change them and every green result recalculates."
ws["A2"].font = Font(italic=True, color="5F6C84")

band(4, "1 · INPUTS")
line(5, "Product", "XAU/USD", "Any product, for your reference", inp=True)
line(6, "Order side", "Buy", "Buy profits when price goes up, Sell when it goes down", inp=True)
line(7, "Simulation direction", "Up", "Up = price hikes, Down = price drops (the Up / Down buttons)", inp=True)
line(8, "Start price", 4520, "Price the order is opened at (admin 'Start price')", "#,##0.00####", inp=True)
line(9, "Amount to invest (margin budget)", 1000, "A number: how much money you put into the order", "#,##0.00", inp=True)
line(10, "Hike value", 15, "15 = +15 % when Hike type is Percent, or +15 in currency when it is Currency", "#,##0.00####", inp=True)
line(11, "Multiplier (leverage)", 100, "Order ticket 'Multiplier' (100 = 1:100)", "0", inp=True)
line(12, "Units per lot", 100, "Platform rule: 1 lot = 100 units", "0", inp=True)
line(13, "Handling fee rate", 0.001, "Platform rule: 0.1% of margin", "0.00%", inp=True)
line(14, "Max lots per order", 100, "Platform limit", "0", inp=True)
line(15, "Min lot step", 0.01, "Platform minimum / step", "0.00", inp=True)
line(16, "Hike type", "Percent", "Percent = % of the start price, Currency = a fixed amount added to / taken off the price", inp=True)

band(17, "2 · WHAT TO TYPE INTO THE ADMIN SIMULATION")
line(18, "Direction button", '=IF(B7="Up","↑ Up","↓ Down")', "Press this on the product's simulation controls", out=True)
line(19, "Start price field", "=B8", "Start price (optional) box", "#,##0.00####", out=True)
line(20, "Hike as percent", '=IF(B16="Percent",B10/100,B10/B8)', "Target price box, unit  %  (move away from the start price)", "0.00%", out=True)
line(21, "Target as price", '=B8*(1+IF(B7="Up",1,-1)*B20)', "…or unit  $  and type this price", "#,##0.00####", out=True)
line(22, "Hike as currency amount", "=ABS(B21-B8)", "Type this as the move when the unit is $", "#,##0.00####;-#,##0.00####", out=True)

band(24, "3 · HOW MANY LOTS TO BUY FOR YOUR AMOUNT")
line(25, "Margin per 1 lot", "=B12*B8/B11", "units × price ÷ multiplier", "#,##0.00", out=True)
line(26, "Cost per 1 lot incl. fee", "=B25*(1+B13)", "margin + handling fee", "#,##0.00", out=True)
line(27, "LOTS TO PURCHASE", "=MIN(B14,ROUNDDOWN(B9/B26/B15,0)*B15)", "Largest lot size your amount covers (rounded down to the lot step)", "0.00", out=True, bold=True)
line(28, "Margin held", "=B27*B25", "", "#,##0.00", out=True)
line(29, "Handling fee", "=B28*B13", "", "#,##0.00", out=True)
line(30, "Total cash needed", "=B28+B29", "must be ≤ your balance", "#,##0.00", out=True)
line(31, "Left over from your amount", "=B9-B30", "", "#,##0.00", out=True)

band(33, "4 · RESULT WHEN THE TARGET IS REACHED")
line(34, "Profit / loss (currency)", '=IF(B6="Buy",1,-1)*(B21-B8)*B27*B12', "(target − start) × lots × units, sign by order side", "#,##0.00;[Red]-#,##0.00", out=True, bold=True)
line(35, "Return on margin (%)", "=IF(B28=0,0,B34/B28)", "profit ÷ margin held - what the account really earns", "0.00%;[Red]-0.00%", out=True)
line(36, "Return on amount invested (%)", "=IF(B9=0,0,B34/B9)", "profit ÷ your amount, after fees: see next row", "0.00%;[Red]-0.00%", out=True)
line(37, "Profit after fee (currency)", "=B34-B29", "", "#,##0.00;[Red]-#,##0.00", out=True)
line(38, "Profit after fee (% of amount)", "=IF(B9=0,0,B37/B9)", "", "0.00%;[Red]-0.00%", out=True)
line(39, "Balance after closing", "=B9+B37", "your amount + profit after fee", "#,##0.00", out=True)
line(40, "Move that wipes out the margin", "=IF(B27=0,0,B28/(B27*B12))", "price move (currency) at which the margin is lost - the stop-out level is reached before this", "#,##0.00####", out=True)

band(42, "5 · REVERSE: I WANT A PROFIT OF …")
line(43, "Profit wanted (currency)", 150, "Type the profit you want in currency", "#,##0.00", inp=True)
line(44, "Lots needed", "=MIN(B14,ROUNDUP((B43+0)/(ABS(B21-B8)*B12)/B15,0)*B15)", "profit ÷ (price move × units), rounded up to the lot step", "0.00", out=True, bold=True)
line(45, "Margin + fee needed", "=B44*B25*(1+B13)", "cash you must have", "#,##0.00", out=True)
line(46, "Profit at those lots", "=B44*ABS(B21-B8)*B12", "", "#,##0.00", out=True)
line(47, "Price move needed for that profit at the lots in section 3", "=IF(B27=0,0,B43/(B27*B12))", "if you keep your amount's lots, the price must move this much (currency)", "#,##0.00####", out=True)
line(48, "…as a % of start price", "=IF(B8=0,0,B47/B8)", "enter this % as the simulation target", "0.00%", out=True)

dv_side = DataValidation(type="list", formula1='"Buy,Sell"', allow_blank=False); ws.add_data_validation(dv_side); dv_side.add("B6")
dv_dir = DataValidation(type="list", formula1='"Up,Down"', allow_blank=False); ws.add_data_validation(dv_dir); dv_dir.add("B7")
dv_type = DataValidation(type="list", formula1='"Percent,Currency"', allow_blank=False); ws.add_data_validation(dv_type); dv_type.add("B16")
for width, col in ((46, "A"), (20, "B"), (70, "C"), (3, "D")):
    ws.column_dimensions[col].width = width
ws.freeze_panes = "A4"

# ---- Scenarios table ----
sc = wb.create_sheet("Scenarios")
sc["A1"] = "Profit at different % moves (uses the lots from the Calculator sheet)"; sc["A1"].font = Font(bold=True, size=13)
heads = ["Move %", "Target price", "Price change", "Profit (currency)", "Profit after fee", "Return on margin", "Return on amount"]
for i, h in enumerate(heads, 1):
    c = sc.cell(row=3, column=i, value=h); c.font = hdr; c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = Alignment(horizontal="center")
moves = [0.01, 0.02, 0.05, 0.10, 0.15, 0.20, 0.25, 0.30, 0.50, 1.00]
for r, m in enumerate(moves, 4):
    sc.cell(r, 1, m).number_format = "0%"; sc.cell(r, 1).fill = PatternFill("solid", fgColor=BLUE_IN); sc.cell(r, 1).font = inp_font
    sc.cell(r, 2, f'=Calculator!$B$8*(1+IF(Calculator!$B$7="Up",1,-1)*A{r})').number_format = "#,##0.00####"
    sc.cell(r, 3, f"=B{r}-Calculator!$B$8").number_format = "#,##0.00####;-#,##0.00####"
    sc.cell(r, 4, f'=IF(Calculator!$B$6="Buy",1,-1)*C{r}*Calculator!$B$27*Calculator!$B$12').number_format = "#,##0.00;[Red]-#,##0.00"
    sc.cell(r, 5, f"=D{r}-Calculator!$B$29").number_format = "#,##0.00;[Red]-#,##0.00"
    sc.cell(r, 6, f"=IF(Calculator!$B$28=0,0,D{r}/Calculator!$B$28)").number_format = "0.0%;[Red]-0.0%"
    sc.cell(r, 7, f"=IF(Calculator!$B$9=0,0,E{r}/Calculator!$B$9)").number_format = "0.0%;[Red]-0.0%"
    for c in range(1, 8): sc.cell(r, c).border = box
for col, w in zip("ABCDEFG", (11, 16, 16, 18, 18, 18, 18)): sc.column_dimensions[col].width = w
sc["A15"] = "Change the blue % values to test any move you like."; sc["A15"].font = Font(italic=True, color="5F6C84")

# ---- Notes ----
nt = wb.create_sheet("How it works")
rows = [
    ("Formulas (same as the trading platform)", ""),
    ("Margin held", "lots × units per lot × price ÷ multiplier"),
    ("Handling fee", "margin × 0.1%"),
    ("Profit / loss", "(exit price − entry price) × lots × units per lot   (negative for a Sell if price rises)"),
    ("Lots for an amount", "amount ÷ (margin per lot × (1 + fee rate)), rounded DOWN to 0.01, max 100"),
    ("Lots for a profit", "profit ÷ (price move × units per lot), rounded UP to 0.01, max 100"),
    ("Target price", "start price × (1 ± move %)   (+ for Up, − for Down)"),
    ("", ""),
    ("Using it with the admin simulation", ""),
    ("1", "Type your amount, start price, hike type (Percent or Currency) and hike value in the blue cells of the Calculator sheet."),
    ("2", "On Admin → Products open the product, put the Start price in, choose Up/Down as shown in section 2."),
    ("3", "Enter the Target as % (section 2) or as a price ($) and press Up/Down."),
    ("4", "Place the order with the LOTS TO PURCHASE shown in section 3."),
    ("", ""),
    ("Note", "The platform also applies a stop-out when margin level falls to 20%, so large leverage with a big adverse move can close the position before the margin is fully lost."),
]
for i, (a, b) in enumerate(rows, 1):
    nt.cell(i, 1, a).font = Font(bold=True); nt.cell(i, 2, b)
nt.column_dimensions["A"].width = 38; nt.column_dimensions["B"].width = 100
nt["A1"].font = Font(bold=True, size=13)
wb.save("tools/Simulation-Calculator.xlsx")
